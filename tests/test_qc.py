import sys, copy, datetime, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl
from core.parsers import parse_jadwal, parse_hrd
from core import engine, qc_layout
from core.depts import DEPTS
from core.render import render
from core.store import Store
A = '/mnt/user-data/attachments/1a0376e9-f1cd-5af9-96bf-99b03536ce1c/'
cfg = DEPTS['QC']
jad = parse_jadwal(A + '5fbf0280-JADWAL_KERJA_OKTOBER_QC.xlsx'); hrd = parse_hrd(A + '1c09f081-23._QC.xlsx')
DEC = {'962': '05D002'}


def r0():
    return engine.compute(jad, hrd['rows'], None, DEC, None, cfg)


def slots(ws):
    return [r for r in range(12, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, int)]


def test_detect_needs_decision_for_unknown_person():
    k = [a for a in engine.detect(jad, hrd['rows'], None, cfg) if a['kind'] == 'jadwal_only']
    assert [a['nik'] for a in k] == ['962'] and '05D002' in k[0]['options'] and k[0]['default'] == '05D002'
    plan, _ = engine.compute(jad, hrd['rows'], None, {'962': '(pilih)'}, None, cfg)
    assert plan['pending'] == ['962'] and '962' in plan['skipped']


def test_r0_structure(tmp_path):
    plan, snap = r0()
    assert len(plan['people']) == 48 and sum(1 for p in plan['people'] if p['regu']) == 37
    bulks, rest = qc_layout.split_bulk(plan, 50)
    assert [(b['kind'], b['date'], len(b['niks'])) for b in bulks] == [('off', 9, 48), ('on', 24, 46)]
    assert sorted(rest) == ['706', '774']
    out = str(tmp_path / 'q.xlsx'); info = render('QC', plan, 0, datetime.date(2026, 10, 9), out)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ['OFF', 'MASUK', 'R0', 'PERUBAHAN REGU']
    assert len(slots(wb['OFF'])) == 48 and len(slots(wb['MASUK'])) == 46 and len(slots(wb['R0'])) == 2
    off = wb['OFF']
    assert {off.cell(r, 10).value for r in slots(off)} == {'9 Oktober 2026'}
    assert off['P3'].value == 'GRESIK, 9 Oktober 2026' and off['P8'].value == 'M. NUR MUSTHOFA'
    assert wb['R0'].cell(12, 16).value == 'OFF' and wb['R0'].cell(12, 19).value == 'DSO2'
    reg = wb['PERUBAHAN REGU']
    assert len(slots(reg)) == 37 and reg['O3'].value == 'GRESIK, 9 Oktober 2026'


def test_r1_r2_qc(tmp_path):
    plan0, snap0 = r0()
    j1 = copy.deepcopy(jad)
    p = next(x for x in j1['people'] if x['nik'] == '93')
    p['days'][15] = 'DS08' if p['days'][15] == 'OFF' else 'OFF'
    q = next(x for x in j1['people'] if x['nik'] == '416')
    q['days'][19] = 'OFF' if q['days'][19] != 'OFF' else 'DS04'
    plan1, snap1 = engine.compute(j1, hrd['rows'], snap0, {}, None, cfg)
    assert sum(len(x['ents']) for x in plan1['people']) == 2 and not any(x['regu'] for x in plan1['people'])
    out = str(tmp_path / 'q1.xlsx'); render('QC', plan1, 1, datetime.date(2026, 10, 20), out)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ['R1'] and len(slots(wb['R1'])) == 2
    assert wb['R1']['P3'].value == 'GRESIK, 20 Oktober 2026'
    j2 = copy.deepcopy(j1); p2 = next(x for x in j2['people'] if x['nik'] == '93'); p2['days'][15] = jad['people'][1]['days'][15]
    plan2, _ = engine.compute(j2, hrd['rows'], snap1, {}, None, cfg)
    ch = [x for x in plan2['people'] if x['ents']]
    assert [x['nik'] for x in ch] == ['93']
    out2 = str(tmp_path / 'q2.xlsx'); assert render('QC', plan2, 2, datetime.date(2026, 10, 25), out2)['sheets'] == ['R2']
    # regu berubah di R1 -> sheet PERUBAHAN REGU muncul
    j3 = copy.deepcopy(jad); next(x for x in j3['people'] if x['nik'] == '416')['pws'] = '05D008'
    plan3, _ = engine.compute(j3, hrd['rows'], snap0, {}, None, cfg)
    info = render('QC', plan3, 1, datetime.date(2026, 10, 20), str(tmp_path / 'q3.xlsx'))
    assert info['sheets'] == ['R1', 'PERUBAHAN REGU'] and info['regu'] == 1


def test_unmapped_person_inherits_team():
    plan, _ = r0()
    teams = [t[:] for t in cfg['teams']]
    teams[1].remove('697')                                   # Indra tidak terdaftar
    groups = qc_layout.order_people(plan['people'], teams)
    assert '697' in [p['nik'] for p in groups[1]]            # ikut tim orang di atasnya (Indra setelah Dani, tim 2)


def test_more_than_three_pages(tmp_path):
    plan, _ = r0()
    teams = [[p['nik'] for p in plan['people'][i:i + 12]] for i in range(0, 48, 12)]   # 4 tim
    out = str(tmp_path / 'q4.xlsx'); render('QC', plan, 0, datetime.date(2026, 10, 9), out, teams)
    wb = openpyxl.load_workbook(out)
    assert len(slots(wb['OFF'])) == 48 and wb['OFF'].max_row >= 208


def test_backup_sheet_names_per_dept(tmp_path, monkeypatch):
    from core import backup
    store = Store(str(tmp_path / 'b.db'))
    plan, snap = r0(); store.save_rev('QC', '2026-10', 0, plan, snap, 'HRD'); store.save_ledger('QC', '2026-10', snap)
    store.save_rev('SK', '2026-10', 0, plan, snap, 'HRD')
    class WS:
        def __init__(s, n, book): s.n, s.book, s.vals = n, book, []
        def clear(s): s.vals = []
        def update(s, values=None, range_name=None): s.vals = values
        def get_all_values(s): return [[str(a), b] for a, b in s.vals]
    class Book:
        def __init__(s): s.w = {}
        def worksheet(s, n):
            if n not in s.w: raise KeyError(n)
            return s.w[n]
        def add_worksheet(s, n, r, c): s.w[n] = WS(n, s); return s.w[n]
    book = Book(); monkeypatch.setattr(backup, '_open', lambda secrets: book)
    backup.push(store, {}, ['SK', 'QC'])
    assert {'data_backup_SK', 'data_backup_QC', 'saldo_hari_SK', 'saldo_hari_QC', 'riwayat_SK', 'riwayat_QC'} <= set(book.w)
    st2 = Store(':memory:'); assert backup.pull(st2, {}, ['SK', 'QC']) == ['SK', 'QC']
    assert st2.revs('QC', '2026-10') == [0] and len(st2.get_ledger('QC', '2026-10')) == 48
    # pulih satu departemen tidak menghapus departemen lain
    st2.save_rev('SK', '2026-11', 0, plan, snap, 'HRD'); backup.pull(st2, {}, ['QC']); assert st2.revs('SK', '2026-11') == [0]


def test_jadwal_output(tmp_path):
    from core.jadwal_out import make_jadwal
    import io
    out = make_jadwal(open(A + '5fbf0280-JADWAL_KERJA_OKTOBER_QC.xlsx', 'rb').read(), 3, datetime.date(2026, 10, 30))
    ws = openpyxl.load_workbook(io.BytesIO(out)).active
    assert ws['B7'].value == 'Revisi : 3' and ws['Z2'].value == 'Gresik, 30 Oktober 2026' and ws['H10'].value == 'DSO4'


def test_r1_same_file_has_no_spurious_regu():
    """R1 memakai HRD tersimpan: keputusan PWS lama (R0) tidak boleh dianggap PWS baru."""
    from core.depts import DEPTS
    A = '/mnt/user-data/attachments/1a0376e9-f1cd-5af9-96bf-99b03536ce1c/'
    j = parse_jadwal(A + '5fbf0280-JADWAL_KERJA_OKTOBER_QC.xlsx'); h = parse_hrd(A + '1c09f081-23._QC.xlsx')
    dec = {'962': '05D002'}
    p0, s0 = engine.compute(j, h['rows'], None, dec, None, DEPTS['QC'])
    p1, _ = engine.compute(j, h['rows'], s0, dec, None, DEPTS['QC'])
    assert any(p['regu'] for p in p0['people']) and not any(p['regu'] for p in p1['people'])
