import openpyxl,warnings,sys;warnings.filterwarnings('ignore')
a=openpyxl.load_workbook(sys.argv[1]);b=openpyxl.load_workbook(sys.argv[2])
print(a.sheetnames,b.sheetnames); bad=0
for x,y in zip(a,b):
    assert x.dimensions==y.dimensions,(x.dimensions,y.dimensions)
    for row in x.iter_rows():
        for c in row:
            if c.value!=y[c.coordinate].value: bad+=1; print(x.title,c.coordinate,repr(c.value),repr(y[c.coordinate].value))
    ma=sorted(map(str,x.merged_cells.ranges)); mb=sorted(map(str,y.merged_cells.ranges))
    if ma!=mb: print('merge diff',x.title,set(ma)^set(mb)); bad+=1
print('DIFFS',bad)
