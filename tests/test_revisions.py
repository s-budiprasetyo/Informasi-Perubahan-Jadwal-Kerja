import sys, copy, datetime, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl
from core.parsers import parse_jadwal, parse_hrd
from core import engine
from core.inf_writer import render_inf
from core.store import Store
A = '/mnt/user-data/attachments/1a0376e9-f1cd-5af9-96bf-99b03536ce1c/'
TPL = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'assets/template_inf.xlsx')
jad = parse_jadwal(A + 'dddeacc0-JADWAL_KERJA_SK.xlsx'); hrd = parse_hrd(A + 'ef819b7b-25._SHUTTLE_KILN_1.xlsx')


def apply(base, ents):
    d = list(base)
    for a, b, ds in ents:
        for x in ds:
            assert d[x - 1] == a, (x, d[x - 1], a); d[x - 1] = b
    return d


def entries_xlsx(path):
    wb = openpyxl.load_workbook(path)
    n = {}
    for ws in wb:
        n[ws.title] = sum(1 for r in range(12, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, int))
    return n


def test_r0_reconstructs():
    plan, snap = engine.compute(jad, hrd['rows'], None, {'409': '2G0301'})
    pat = engine.patterns(hrd['rows'])
    assert len(plan['people']) == 28 and not plan['skipped']
    for p, sp in zip(plan['people'], snap['people']):
        s = hrd['rows'][p['nik']]
        if p['regu']:
            E = p['regu']['tgl']; base = pat[p['pws_lama']][:E - 1] + pat[p['pws']][E - 1:]
        else:
            base = pat[p['pws_lama'] if p['exempt'] else p['pws']] if not p['exempt'] else pat['05D004']
        assert apply(base, p['ents']) == sp['days'], p['nama']


def mutate(j, changes):
    j2 = copy.deepcopy(j)
    for p in j2['people']:
        for nik, (day, code) in changes.items():
            if p['nik'] == nik: p['days'][day - 1] = code
    return j2


def test_r1_r2_chain(tmp_path):
    plan0, snap0 = engine.compute(jad, hrd['rows'], None, {'409': '2G0301'})
    sulton = next(p for p in jad['people'] if p['nama'].startswith('SULTON') or p['nik'] == '1047')['nik']
    qcnik = next(p['nik'] for p in jad['people'] if p['seksi'] == 'QCSK')
    j1 = mutate(jad, {sulton: (3, 'OFF' if jad['people'][0]['days'][2] != 'OFF' else 'DS04')})
    # ubah satu orang QC & satu orang shift
    shiftp = next(p for p in jad['people'] if p['pws'] in ('2G0301',) and p['seksi'] != 'QCSK')
    cur = shiftp['days'][4]
    j1 = mutate(j1, {shiftp['nik']: (5, 'S005' if cur != 'S005' else 'P004')})
    plan1, snap1 = engine.compute(j1, hrd['rows'], snap0, {})
    changed = [p for p in plan1['people'] if p['ents']]
    assert 1 <= len(changed) <= 3, [(p['nama'], p['ents']) for p in changed]
    for p in plan1['people']:
        base = next(x for x in snap0['people'] if x['nik'] == p['nik'])['days']
        assert apply(base, p['ents']) == next(x for x in snap1['people'] if x['nik'] == p['nik'])['days']
    out1 = str(tmp_path / 'r1.xlsx')
    info = render_inf(plan1, 1, datetime.date(2026, 10, 20), TPL, out1)
    assert info['sheets'] == ['SATRIO R1', 'SULTON R1', 'PERUBAHAN REGU R1']
    n = entries_xlsx(out1)
    assert n['SULTON R1'] + n['SATRIO R1'] == sum(len(p['ents']) for p in plan1['people'])
    # R2 dari R1: perubahan lagi pada orang yang sama hanya relatif R1
    j2 = mutate(j1, {shiftp['nik']: (5, cur)})          # kembali ke semula
    plan2, snap2 = engine.compute(j2, hrd['rows'], snap1, {})
    ch2 = [p for p in plan2['people'] if p['ents']]
    assert len(ch2) == 1 and ch2[0]['nik'] == shiftp['nik']
    out2 = str(tmp_path / 'r2.xlsx')
    assert render_inf(plan2, 2, datetime.date(2026, 10, 25), TPL, out2)['sheets'][1] == 'SULTON R2'
    # tidak ada perubahan sama sekali -> file tetap valid
    plan3, _ = engine.compute(j2, hrd['rows'], snap2, {})
    assert not any(p['ents'] or p['regu'] for p in plan3['people'])
    render_inf(plan3, 3, datetime.date(2026, 10, 26), TPL, str(tmp_path / 'r3.xlsx'))
    openpyxl.load_workbook(str(tmp_path / 'r3.xlsx'))


def test_regu_change_in_r1(tmp_path):
    plan0, snap0 = engine.compute(jad, hrd['rows'], None, {'409': '2G0301'})
    j1 = copy.deepcopy(jad)
    p = next(x for x in j1['people'] if x['pws'] == '05D004')
    p['pws'] = '05D002'
    plan1, _ = engine.compute(j1, hrd['rows'], snap0, {})
    r = next(x for x in plan1['people'] if x['nik'] == p['nik'])
    assert r['regu'] and r['regu']['lama'] == '05D004' and r['regu']['baru'] == '05D002'
    render_inf(plan1, 1, datetime.date(2026, 10, 20), TPL, str(tmp_path / 'x.xlsx'))


def test_ledger_carry(tmp_path):
    st = Store(str(tmp_path / 'a.db'))
    plan0, snap0 = engine.compute(jad, hrd['rows'], None, {'409': '2G0301'})
    snap0['people'][0]['kurang'] = 4
    st.save_ledger('SK', '2026-10', snap0)
    nov = copy.deepcopy(snap0); nov['people'][0]['kurang'] = -1
    st.save_ledger('SK', '2026-11', nov)
    row = next(r for r in st.get_ledger('SK', '2026-11') if r['nik'] == snap0['people'][0]['nik'])
    assert row['saldo_awal'] == 4 and row['saldo_akhir'] == 3


def test_store_revs_and_backup(tmp_path):
    st = Store(str(tmp_path / 'b.db'))
    assert st.next_rev('SK', '2026-10') == 0
    plan0, snap0 = engine.compute(jad, hrd['rows'], None, {'409': '2G0301'})
    st.save_rev('SK', '2026-10', 0, plan0, snap0, 'HRD'); st.save_rev('SK', '2026-10', 1, plan0, snap0, 'R0')
    assert st.next_rev('SK', '2026-10') == 2
    st.save_rev('SK', '2026-10', 0, plan0, snap0, 'HRD')       # timpa R0 -> R1 gugur
    assert st.revs('SK', '2026-10') == [0]
    data = st.export_all(); st2 = Store(':memory:'); st2.import_all(data)
    assert st2.get_rev('SK', '2026-10', 0)['plan']['ym'] == '2026-10'


def test_delete_from_recomputes_ledger(tmp_path):
    st = Store(str(tmp_path / 'd.db'))
    plan0, snap0 = engine.compute(jad, hrd['rows'], None, {'409': '2G0301'})
    st.save_rev('SK', '2026-10', 0, plan0, snap0, 'HRD', 'j.xlsx', 'h.xlsx')
    st.save_ledger('SK', '2026-10', snap0)
    st.set_setting('SK', 'holidays:2026-10', [17])
    plan1, snap1 = engine.compute(jad, hrd['rows'], snap0, {'409': '2G0301'})
    st.save_rev('SK', '2026-10', 1, plan1, snap1, 'R0', 'j.xlsx', '')
    st.delete_from('SK', '2026-10', 1)
    assert st.revs('SK', '2026-10') == [0]
    st.delete_from('SK', '2026-10', 0)
    assert st.revs('SK', '2026-10') == [] and not st.ledger_months('SK')
    assert st.get_setting('SK', 'holidays:2026-10', None) is None


def test_holiday_suppresses_change():
    d0 = {'409': '2G0301'}
    p0, _ = engine.compute(jad, hrd['rows'], None, d0)
    cand = [(p['nik'], x) for p in p0['people'] for a, b, ds in p['ents'] if b == 'OFF' for x in ds]
    assert cand
    nik, day = cand[0]
    p1, s1 = engine.compute(jad, hrd['rows'], None, d0, None, None, [day])
    q = next(p for p in p1['people'] if p['nik'] == nik)
    assert all(day not in ds for a, b, ds in q['ents'])
    assert p1['holidays'] == [day]


def test_potong_cuti_chain_and_recap(tmp_path):
    st = Store(str(tmp_path / 'c.db'))
    snap = lambda k: {'people': [dict(nik='260', nama='A', target=22, worked=22 - k, kurang=k)]}
    st.save_ledger('SK', '2026-10', snap(4)); st.save_ledger('SK', '2026-11', snap(-1))
    st.add_cuti('SK', '2026-10', '260', 'A', '2026-10-20', 1, 'cuti')
    st.add_cuti('SK', '2026-10', '260', 'A', '2026-10-25', 2, '')
    log = st.cuti_log('SK')
    assert [r['sisa'] for r in sorted(log, key=lambda r: r['tgl'])] == [3, 1]
    assert st.get_ledger('SK', '2026-10')[0]['saldo_akhir'] == 1
    assert st.get_ledger('SK', '2026-11')[0]['saldo_awal'] == 1 and st.get_ledger('SK', '2026-11')[0]['saldo_akhir'] == 0
    st.delete_cuti('SK', log[0]['id'])
    assert st.get_ledger('SK', '2026-11')[0]['saldo_akhir'] == 2
    # ekspor/impor ikut membawa potong cuti
    st2 = Store(':memory:'); st2.import_all(st.export_all('SK'), 'SK')
    assert len(st2.cuti_log('SK')) == 1 and st2.get_ledger('SK', '2026-11')[0]['saldo_akhir'] == 2


def test_jadwal_recap_fills_cuti_column():
    from core.jadwal_out import make_jadwal, _locate, _recap_map
    d = open(A + 'dddeacc0-JADWAL_KERJA_SK.xlsx', 'rb').read()
    t, _, _ = _locate(d); c, o, rows, lb = _recap_map(d, t)
    out = make_jadwal(d, 1, datetime.date(2026, 10, 31), {'260': {'cuti': 1.5, 'open': 3, 'lebih': 2}})
    ws = openpyxl.load_workbook(__import__('io').BytesIO(out))[t]
    assert ws.cell(rows['260'], c).value == 1.5 and ws.cell(rows['260'], o).value == 3 and ws.cell(rows['260'], lb).value == 2


def test_read_balances_from_jadwal():
    from core.sheetbal import read_balances
    q = read_balances(open(A + '5fbf0280-JADWAL_KERJA_OKTOBER_QC.xlsx', 'rb').read(), 'qc')
    assert q['23'] == dict(awal=2, baru=0, lebih=0, cuti=0, akhir=2) and q['416']['akhir'] == 2 and q['136']['akhir'] == 0
    s = read_balances(open(A + 'dddeacc0-JADWAL_KERJA_SK.xlsx', 'rb').read(), 'sk')
    assert s['1047']['baru'] == 1 and s['1047']['akhir'] == 1


def test_ledger_uses_sheet_values(tmp_path):
    st = Store(str(tmp_path / 'b.db'))
    st.save_ledger('QC', '2026-10', {'people': [dict(nik='23', nama='S', kurang=0, awal=2.0, lebih=0)]})
    assert st.get_ledger('QC', '2026-10')[0]['saldo_akhir'] == 2
    st.save_ledger('QC', '2026-11', {'people': [dict(nik='23', nama='S', kurang=0, lebih=1.0)]})
    r = st.get_ledger('QC', '2026-11')[0]
    assert r['saldo_awal'] == 2 and r['saldo_akhir'] == 1
    assert st.opening_known('QC', '2026-12', '23') == 1 and st.opening_known('QC', '2026-10', '23') is None


def test_matches_manual_sk(tmp_path):
    """Hasil SK harus sama isinya dengan form buatan manual (Oktober): termasuk Dio (ada di HRD -> tanggal 1-7 tetap dicatat)."""
    import re
    plan, _ = engine.compute(jad, hrd['rows'], None, {'409': '2G0301'})
    out = str(tmp_path / 'sk.xlsx'); render_inf(plan, 0, datetime.date(2026, 10, 9), TPL, out)
    man = os.path.join(os.path.dirname(__file__), 'fixtures', 'manual_sk_oktober.xlsx')

    def dump(path):
        wb = openpyxl.load_workbook(path); res = {}
        for ws in wb:
            res[ws.title] = [[[re.sub(r'\s+', ' ', str(ws.cell(r + k, c).value).strip()) for c in range(1, 30) if ws.cell(r + k, c).value is not None] for k in (0, 1)]
                             for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, int)]
        return res
    a, b = dump(out), dump(man)
    assert a == b
    dio = next(r for r in a['SULTON'] if r[0][1] == '871')
    assert '1, 2, 5, 6, 7, 14, 20 & 26 Oktober 2026' in dio[0]


def test_new_person_not_in_hrd_skips_days_before_start():
    """Orang yang TIDAK ada di HRD (karyawan baru): libur sebelum tanggal masuk bukan perubahan."""
    j = copy.deepcopy(jad); hr = copy.deepcopy(hrd['rows'])
    p = next(x for x in j['people'] if x['nik'] == '871'); del hr['871']
    pl, _ = engine.compute(j, hr, None, {'409': '2G0301', '871': '05D004'})
    q = next(x for x in pl['people'] if x['nik'] == '871')
    assert all(not any(d < 8 for d in ds) for a_, b_, ds in q['ents'])


def test_sk_matches_manual_form():
    """R0 SK harus sama isinya dengan form manual (semua sheet, semua baris) - termasuk karyawan lama yang libur tgl 1-7."""
    import re
    from core.depts import DEPTS
    from core.render import render
    import tempfile
    plan, _ = engine.compute(jad, hrd['rows'], None, {'409': '2G0301', '1074': 'abaikan', '1075': 'abaikan'}, None, DEPTS['SK'])
    out = tempfile.mktemp(suffix='.xlsx'); render('SK', plan, 0, datetime.date(2026, 10, 9), out)

    def dump(p):
        wb = openpyxl.load_workbook(p)
        return {ws.title: [[[re.sub(r'\s+', ' ', str(ws.cell(r + k, c).value).strip()) for c in range(1, 30) if ws.cell(r + k, c).value is not None]
                            for k in (0, 1)] for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, int)] for ws in wb}
    assert dump(out) == dump(os.path.join(os.path.dirname(__file__), 'fixtures', 'manual_sk_oktober.xlsx'))
