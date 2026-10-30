"""Pembaca file: Jadwal Kerja SK (lapangan) dan HRD '25. Shuttle Kiln 1'."""
import re, warnings
import openpyxl
from .consts import norm, parse_month_text, month_len

warnings.filterwarnings('ignore')


def _wb(f):
    return openpyxl.load_workbook(f, data_only=False)


def parse_jadwal(f):
    """-> dict(ym, target, people=[dict(no,nama,nik,item,seksi,pws,days[31])])"""
    wb = _wb(f)
    ws = next((s for s in wb if 'JADWAL KERJA' in str(s['A3'].value or '').upper()), wb.active)
    ym = parse_month_text(ws['A3'].value)
    if ym is None:
        raise ValueError('Bulan tidak terbaca dari sel A3 ("JADWAL KERJA KARYAWAN BULAN ...").')
    nd = month_len(ym)
    target = ws['AQ8'].value if isinstance(ws['AQ8'].value, (int, float)) else 22
    rev_txt = str(ws['A7'].value or '')
    people = []
    for row in range(10, ws.max_row + 1):
        nama = ws.cell(row, 2).value
        nik = str(ws.cell(row, 3).value or '').strip()
        if not nama or not nik.isdigit():
            continue
        days = [norm(ws.cell(row, c).value) for c in range(8, 8 + nd)]
        people.append(dict(no=ws.cell(row, 1).value, nama=str(nama).strip(), nik=str(ws.cell(row, 3).value).strip(),
                           item=ws.cell(row, 4).value, seksi=str(ws.cell(row, 5).value or '').strip().upper(),
                           pws=str(ws.cell(row, 7).value or '').strip(), days=days))
    return dict(ym=ym, target=target, rev_text=rev_txt, people=people)


def parse_hrd(f):
    """-> dict(ym, rows={nik: dict(nik,dept,nama,pws,days,total)})"""
    wb = _wb(f)
    ws = wb.active
    ym = parse_month_text(ws['B2'].value) or parse_month_text(ws['A2'].value)
    nd = month_len(ym) if ym else 31
    hdr = 4
    tot_col = None
    for c in range(1, ws.max_column + 1):
        if 'TOTAL' in str(ws.cell(hdr, c).value or '').upper():
            tot_col = c
    rows = {}
    for r in range(hdr + 1, ws.max_row + 1):
        nik = ws.cell(r, 1).value
        if nik is None or str(nik).strip() == '':
            continue
        nik = str(nik).strip()
        total = ws.cell(r, tot_col).value if tot_col else None
        rows[nik] = dict(nik=nik, dept=ws.cell(r, 2).value, nama=str(ws.cell(r, 3).value or '').strip(),
                         pws=str(ws.cell(r, 4).value or '').strip(),
                         days=[norm(ws.cell(r, c).value) for c in range(5, 5 + nd)],
                         total=total if isinstance(total, (int, float)) else None)
    return dict(ym=ym, rows=rows)
