"""Berkas 'Jadwal Kerja' hasil render: salinan file Jadwal Kerja yang diunggah (format & gambar asli utuh),
dengan 'Revisi : n' dan tanggal 'Gresik, ...' diperbarui."""
import io, re, zipfile
from openpyxl.utils import get_column_letter, column_index_from_string
import openpyxl
from lxml import etree
from .consts import BULAN_ID

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
RNS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def _locate(data):
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    for ws in wb.worksheets:
        rev = date = None
        for row in ws.iter_rows(min_row=1, max_row=12):
            for c in row:
                v = str(c.value or '').strip()
                if v.lower().startswith('revisi') and rev is None: rev = c.coordinate
                if v.lower().startswith('gresik,') and date is None: date = c.coordinate
        if rev or date:
            return ws.title, rev, date
    raise ValueError('Sel "Revisi : n" / "Gresik, ..." tidak ditemukan di Jadwal Kerja.')


def _recap_map(data, title):
    """Cari kolom 'POTONG CUTI'/'POT CUTI' & kolom saldo bulan lalu ('KRNG. HARI ...'/'LAMA'), serta baris tiap NIK."""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    ws = wb[title]; cuti = open_ = lebih = None; nik_col = hdr = None
    for row in ws.iter_rows(min_row=1, max_row=14):
        for c in row:
            v = str(c.value or '').strip().upper()
            if cuti is None and v in ('POTONG CUTI', 'POT CUTI'): cuti = c.column
            if open_ is None and (v.startswith('KRNG. HARI') or v.startswith('KRNG HARI') or v == 'LAMA'): open_ = c.column
            if lebih is None and v.startswith('GANTI HARI'): lebih = c.column
            if nik_col is None and 'PRN SAP' in v: nik_col, hdr = c.column, c.row
    rows = {}
    if nik_col:
        for r in ws.iter_rows(min_row=hdr + 1, min_col=nik_col, max_col=nik_col):
            c = r[0]; v = re.sub(r'\D', '', str(c.value or ''))
            if v and v not in rows: rows[v] = c.row
    return cuti, open_, rows, lebih


def _setnum(root, ref, val):
    """Isi angka di sel; buat sel bila belum ada (urut kolom). Mengembalikan True bila diisi."""
    col = re.match(r'[A-Z]+', ref).group(); rown = int(ref[len(col):]); ci = column_index_from_string(col)
    for r in root.iter('{%s}row' % NS):
        if int(r.get('r')) != rown: continue
        cell = None
        for c in r.findall('{%s}c' % NS):
            if c.get('r') == ref: cell = c; break
        if cell is None:
            cell = etree.Element('{%s}c' % NS); cell.set('r', ref)
            pos = len(r)
            for i, c in enumerate(r.findall('{%s}c' % NS)):
                if column_index_from_string(re.match(r'[A-Z]+', c.get('r')).group()) > ci: pos = list(r).index(c); break
            r.insert(pos, cell)
        for ch in list(cell): cell.remove(ch)
        if 't' in cell.attrib: del cell.attrib['t']
        v = etree.SubElement(cell, '{%s}v' % NS); v.text = repr(float(val)) if float(val) != int(val) else str(int(val))
        return True
    return False


def _cell_empty(root, ref):
    for c in root.iter('{%s}c' % NS):
        if c.get('r') == ref: return c.find('{%s}v' % NS) is None and c.find('{%s}is' % NS) is None and c.find('{%s}f' % NS) is None
    return True


def _setcell(root, ref, text):
    for c in root.iter('{%s}c' % NS):
        if c.get('r') == ref:
            for ch in list(c): c.remove(ch)
            c.set('t', 'inlineStr')
            t = etree.SubElement(etree.SubElement(c, '{%s}is' % NS), '{%s}t' % NS); t.text = text
            t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            return True
    return False


def make_jadwal(data, rev, gen_date, recap=None):
    """recap: {nik: {'open': saldo awal, 'lebih': lebih hari, 'cuti': potong cuti}} -> MENIMPA kolom saldo awal / ganti hari / potong cuti
    di Jadwal Kerja (hanya dipakai bila pengguna memilih data historis aplikasi, atau saat unduh ulang dari Riwayat)."""
    title, rev_ref, date_ref = _locate(data)
    zin = zipfile.ZipFile(io.BytesIO(data))
    wbx = etree.fromstring(zin.read('xl/workbook.xml'))
    rid = next(s.get('{%s}id' % RNS) for s in wbx.iter('{%s}sheet' % NS) if s.get('name') == title)
    rels = etree.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
    target = next(r.get('Target') for r in rels if r.get('Id') == rid)
    part = 'xl/' + target.lstrip('/').replace('xl/', '', 1) if not target.startswith('/') else target.lstrip('/')
    root = etree.fromstring(zin.read(part))
    if rev_ref: _setcell(root, rev_ref, f'Revisi : {rev}')
    if date_ref: _setcell(root, date_ref, f'Gresik, {gen_date.day} {BULAN_ID[gen_date.month - 1]} {gen_date.year}')
    if recap:
        cuti_c, open_c, rows, lebih_c = _recap_map(data, title)
        for nik, rc in recap.items():
            r = rows.get(str(nik))
            if not r: continue
            for c, k in ((cuti_c, 'cuti'), (open_c, 'open'), (lebih_c, 'lebih')):
                if c and rc.get(k) is not None: _setnum(root, f'{get_column_letter(c)}{r}', rc[k])
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for item in zin.infolist():
            raw = zin.read(item.filename)
            if item.filename == part:
                raw = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
            if item.filename == 'xl/workbook.xml' and b'fullCalcOnLoad' not in raw:
                raw = re.sub(rb'<calcPr([^>]*?)/>', rb'<calcPr\1 fullCalcOnLoad="1"/>', raw)
            z.writestr(item, raw)
    return out.getvalue()
