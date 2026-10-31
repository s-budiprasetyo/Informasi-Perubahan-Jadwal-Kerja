"""Baca saldo kurang/lebih hari dari file Jadwal Kerja yang diunggah (dibaca dari nilai yang tersimpan di Excel).

Kolom dicari lewat judulnya (bukan huruf kolom), jadi cocok untuk tata letak SK dan QC:
  saldo awal : 'LAMA' (QC) / 'KRNG. HARI ...' (SK)
  kurang baru: 'BARU' (QC)  | SK memakai kolom 'Kurang Hari' (AN)
  lebih hari : 'GANTI HARI' (QC AQ) / 'ganti hari ...' (SK)
  potong cuti: 'POT CUTI' (QC AS) / 'POTONG CUTI' (SK)
  sisa       : 'SISA' (QC AW, sama dengan AN)
saldo akhir = saldo awal + kurang baru - lebih hari - potong cuti  (QC memakai nilai kolom SISA bila ada)
"""
import io, re
import openpyxl


def _num(v):
    try: return float(v)
    except (TypeError, ValueError): return 0.0


def read_balances(data, kind):
    """-> {nik: dict(awal, baru, lebih, cuti, akhir)}; {} bila kolom saldo tidak ditemukan."""
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    for ws in wb.worksheets:
        col = {}; nik_col = hdr = None
        for row in ws.iter_rows(min_row=1, max_row=14):
            for c in row:
                v = str(c.value or '').strip().upper()
                if not v: continue
                if 'PRN SAP' in v and nik_col is None: nik_col, hdr = c.column, c.row
                elif v == 'KURANG HARI' and 'an' not in col: col['an'] = c.column
                elif (v == 'LAMA' or v.startswith('KRNG. HARI') or v.startswith('KRNG HARI')) and 'awal' not in col: col['awal'] = c.column
                elif v == 'BARU' and 'baru' not in col: col['baru'] = c.column
                elif v.startswith('GANTI HARI') and 'lebih' not in col: col['lebih'] = c.column
                elif v in ('POTONG CUTI', 'POT CUTI') and 'cuti' not in col: col['cuti'] = c.column
                elif v == 'SISA' and 'sisa' not in col: col['sisa'] = c.column
        if not nik_col or not ({'awal', 'lebih', 'cuti'} & set(col)): continue
        out = {}
        for r in range(hdr + 1, ws.max_row + 1):
            nik = re.sub(r'\D', '', str(ws.cell(r, nik_col).value or ''))
            if not nik or nik in out: continue
            g = lambda k: _num(ws.cell(r, col[k]).value) if k in col else 0.0
            awal, lebih, cuti = g('awal'), g('lebih'), g('cuti')
            baru = g('baru') if kind == 'qc' else g('an')
            akhir = awal + baru - lebih - cuti
            if kind == 'qc' and 'sisa' in col and ws.cell(r, col['sisa']).value is not None: akhir = g('sisa')
            out[nik] = dict(awal=awal, baru=baru, lebih=lebih, cuti=cuti, akhir=akhir)
        return out
    return {}
